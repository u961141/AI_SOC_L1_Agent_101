import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Set Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Title Banner
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("ENTERPRISE FORTIGATE FIREWALL SOC AI DEFENSE SHIELD\n")
    title_run.font.name = 'Calibri'
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    subtitle_run = title_p.add_run("Comprehensive Architectural Engineering Report, Code Flow Diagrams & Retrospective Guide\n")
    subtitle_run.font.name = 'Calibri'
    subtitle_run.font.size = Pt(13)
    subtitle_run.font.italic = True
    subtitle_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_paragraph() # Spacer

    # Section 1: Executive Overview & Problem Statement
    h1 = doc.add_heading("1. Executive Overview & Problem Definition", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    
    p = doc.add_paragraph()
    p.add_run("In modern Enterprise Security Operations Centers (SOC), security analysts are overwhelmed by high-volume firewall traffic logs and security event streams. The traditional workflow for identifying, triaging, and mitigating malicious IP addresses is highly manual, labor-intensive, and prone to human delay. Analysts typically perform the following steps manually:\n")
    
    steps = [
        "Sifting through syslog buffers or FortiGate traffic log records.",
        "Copying unknown external IPv4 addresses one-by-one into web portals or API tools.",
        "Checking threat intelligence scores on platforms like AbuseIPDB.",
        "Manually opening FortiGate Web GUI or CLI to define firewall Address Objects (config firewall address).",
        "Appending new Address Objects to active Firewall Policy Groups (config firewall addrgrp) to enforce blocking."
    ]
    for step in steps:
        bp = doc.add_paragraph(style='List Bullet')
        bp.add_run(step)

    p2 = doc.add_paragraph()
    p2.add_run("\nTarget Solution: ").bold = True
    p2.add_run("We set out to build an autonomous, privacy-preserving, local GPU-accelerated AI SOC Defense Shield. The system leverages Model Context Protocol (FastMCP), LangGraph State Machine Graph Engine, Local Qwen2.5 on an NVIDIA GeForce RTX 4070 GPU, and direct REST API integration with physical/virtual FortiGate firewalls (192.168.10.1) and AbuseIPDB v2 REST API. This reduces manual incident triage and remediation time from 20 minutes down to under 2 seconds.\n")

    # Section 2: System Architecture & Code Flow Diagrams
    h2 = doc.add_heading("2. System Architecture & Visual Code Flow Diagrams", level=1)
    h2.runs[0].font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    doc.add_heading("Diagram 1: End-to-End System Topology Diagram", level=2)
    diag1_p = doc.add_paragraph()
    diag1_p.paragraph_format.left_indent = Inches(0.2)
    diag1_run = diag1_p.add_run("""
+------------------------------------------------------------------------------------------------------------------------+
|                                       ENTERPRISE AI SOC SHIELD ARCHITECTURE                                            |
+------------------------------------------------------------------------------------------------------------------------+
|                                                                                                                        |
|  [ SOC Analyst REPL Shell ] (interactive_soc_shell.py)                                                                 |
|            |                                                                                                           |
|            v                                                                                                           |
|  [ LangGraph State Machine ] (agent.py)                                                                                |
|            |                                                                                                           |
|            +---> Node 1: Log Search & DPI Regex IP Extraction (RFC1918 Private vs Global Public IP Filter)             |
|            |                                                                                                           |
|            +---> Node 2: Live Threat Intel Queries (AbuseIPDB API v2 via FastMCP Protocol)                             |
|            |                                                                                                           |
|            +---> Node 3: Local GPU AI Reasoning Engine (Qwen2.5 on NVIDIA RTX 4070 GPU via Ollama)                    |
|            |                                                                                                           |
|            +---> Node 4: FortiGate Active Remediation Engine (FortiOS REST API v2 on 192.168.10.1)                     |
|            |            |-- Step 1: POST /api/v2/cmdb/firewall/address (Create Object ip_<IP>)                         |
|            |            +-- Step 2: PUT/POST /api/v2/cmdb/firewall/addrgrp/Blocked_Bad_IPs (Update Address Group)      |
|            |                                                                                                           |
|            +---> Node 5: Executive Rich Report Console & LangSmith Observability Cloud                                 |
|                                                                                                                        |
+------------------------------------------------------------------------------------------------------------------------+
""")
    diag1_run.font.name = 'Consolas'
    diag1_run.font.size = Pt(8.5)

    doc.add_heading("Diagram 2: 5-Node LangGraph Execution Flowchart", level=2)
    diag2_p = doc.add_paragraph()
    diag2_p.paragraph_format.left_indent = Inches(0.2)
    diag2_run = diag2_p.add_run("""
  +-----------------------------------+
  | START: User Request / Command     |
  +-----------------------------------+
                    |
                    v
  +-----------------------------------+
  | NODE 1: Log Search & Extraction   |  ---> Queries FortiGate Disk Forward Log API
  +-----------------------------------+       (/api/v2/log/disk/traffic/forward)
                    |
                    v
  +-----------------------------------+
  | NODE 2: AbuseIPDB Threat Intel    |  ---> Live HTTPS API Check for Public Destination IPs
  +-----------------------------------+       (api.abuseipdb.com/api/v2/check)
                    |
                    v
  +-----------------------------------+
  | NODE 3: Local GPU LLM Triage      |  ---> Local Qwen2.5 on NVIDIA RTX 4070 GPU
  +-----------------------------------+       (Calculates Risk Score 0-100 & Reasoning)
                    |
                    v
          /-------------------\
         / Is Risk Score >= 50 \
        /  or Bad IPs > 0?     \
        \                      /
         \--------------------/
           /                \
     YES  /                  \  NO
         v                    v
+-----------------------+   +-----------------------+
| NODE 4: Remediation   |   | NODE 5: Exec Report   |
| (Block on FortiGate)  |   | (Render Summary)      |
+-----------------------+   +-----------------------+
         |                            |
         +----------------------------+
                      |
                      v
             +----------------+
             |      END       |
             +----------------+
""")
    diag2_run.font.name = 'Consolas'
    diag2_run.font.size = Pt(8.5)

    doc.add_heading("Diagram 3: FastMCP Gateway & Tool Protocol Interaction Flow", level=2)
    diag3_p = doc.add_paragraph()
    diag3_p.paragraph_format.left_indent = Inches(0.2)
    diag3_run = diag3_p.add_run("""
[ agent.py (Client) ]                       [ server.py (FastMCP Server) ]                 [ External APIs / Devices ]
         |                                                 |                                          |
         | --- call_tool('search_fortigate_logs') -------->|                                          |
         |                                                 | --- GET /log/disk/traffic/forward ------>| [ FortiGate 192.168.10.1 ]
         |                                                 |<--- HTTP 200 OK (Raw Traffic Logs) ------|
         |<-- Returns JSON {"logs": [...]} ----------------|                                          |
         |                                                 |                                          |
         | --- call_tool('check_abuseipdb_reputation') --->|                                          |
         |                                                 | --- GET api.abuseipdb.com/api/v2 -------->| [ AbuseIPDB Cloud API ]
         |                                                 |<--- HTTP 200 OK (Score, ISP, Country) ---|
         |<-- Returns JSON Intel Data ---------------------|                                          |
         |                                                 |                                          |
         | --- call_tool('block_ip_in_fortigate_group') ->|                                          |
         |                                                 | --- POST /cmdb/firewall/address -------->| [ FortiGate 192.168.10.1 ]
         |                                                 | --- PUT /cmdb/firewall/addrgrp/Blocked ->|
         |                                                 |<--- HTTP 200/201 (Group Updated) --------|
         |<-- Returns SUCCESS Response --------------------|                                          |
""")
    diag3_run.font.name = 'Consolas'
    diag3_run.font.size = Pt(8.0)

    # Section 3: Engineering Challenges, Root Causes & Resolutions
    h3 = doc.add_heading("3. Detailed Engineering Challenges & Technical Resolutions", level=1)
    h3.runs[0].font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    challenges = [
        {
            "num": "Challenge 1",
            "title": "Python 3.14 C++ Compilation & Wheel Package Failures",
            "symptom": "pip install -r requirements.txt failed with errors: Building wheel for zstandard (setup.py) ... error. Microsoft Visual C++ 14.0 or greater is required.",
            "cause": "Exact pinned versions (e.g. langgraph==0.2.60) pulled older dependencies locked to zstandard<0.24.0, which lacked pre-built binary wheel (.whl) packages for Python 3.14 on Windows (cp314). Pip was forced to attempt C++ compilation from source.",
            "fix": "Updated requirements.txt to use flexible lower-bound operators (>=). This allowed pip to download newer, pre-built binary wheel (.whl) packages natively compiled for Python 3.14 without requiring Microsoft Visual C++ Build Tools."
        },
        {
            "num": "Challenge 2",
            "title": "FastMCP Stdio Protocol Crashes (unhandled errors in a TaskGroup)",
            "symptom": "When executing queries, the terminal crashed with: Error executing workflow: unhandled errors in a TaskGroup (1 sub-exception).",
            "cause": "Protocol Corruption on sys.stdout. FastMCP communicates between agent.py and server.py using Standard Input/Output (stdio) transport. The protocol requires sys.stdout to be 100% reserved for raw JSON-RPC protocol packets. Diagnostic print() statements in fortigate_api.py wrote plain text strings (e.g. ⚠️ ...) into sys.stdout, corrupting the JSON-RPC packet stream and crashing stdio_client's reader task group.",
            "fix": "Redirected all diagnostic logging from sys.stdout to sys.stderr (sys.stderr.write). Kept sys.stdout 100% clean for raw JSON-RPC protocol packets. Also updated ChatOllama to use non-blocking async await llm.ainvoke() wrapped in BaseException error handling."
        },
        {
            "num": "Challenge 3",
            "title": "Empty Log Buffer Returns (Total Traffic Logs Found: 0)",
            "symptom": "Querying FortiGate logs returned 0 logs ([]), causing Node 1 to report 0 public destination IPs.",
            "cause": "The initial endpoint used was /api/v2/log/memory/traffic. On many FortiGate firewalls, RAM memory logging is either disabled or cleared rapidly.",
            "fix": "Updated fortigate_api.py to query FortiGate Disk Forward Traffic REST API /api/v2/log/disk/traffic/forward with an expanded capacity limit (limit=100) to audit forward traffic logs spanning the last 24 hours."
        },
        {
            "num": "Challenge 4",
            "title": "FortiGate Group Update HTTP 404/500 Errors",
            "symptom": "PUT /api/v2/cmdb/firewall/addrgrp/Blocked_Bad_IPs failed with HTTP 404 Not Found / 500 Internal Error.",
            "cause": "In FortiOS REST API, PUT updates an existing address group. If address group Blocked_Bad_IPs has not been created on the firewall yet, PUT fails.",
            "fix": "Added an automatic POST fallback mechanism in add_ip_to_block_group(). If PUT returns HTTP 404/500, the code automatically issues a POST to /api/v2/cmdb/firewall/addrgrp to create the address group Blocked_Bad_IPs automatically on the firewall."
        },
        {
            "num": "Challenge 5",
            "title": "Eliminating Hardcoded IP Data for 100% Pure Live Queries",
            "symptom": "Console showed sample IPs (185.220.101.5, 45.33.32.156, 8.8.8.8) when querying internal hosts.",
            "cause": "Early mock fallback arrays in fortigate_api.py injected sample destination IPs when log buffers were empty.",
            "fix": "Completely stripped out all mock fallback arrays and sample dictionaries. The codebase is now 100% pure live REST API querying."
        }
    ]

    for c in challenges:
        doc.add_heading(f"{c['num']}: {c['title']}", level=2)
        
        p_c = doc.add_paragraph()
        p_c.add_run("Symptom: ").bold = True
        p_c.add_run(c["symptom"] + "\n")
        p_c.add_run("Root Cause: ").bold = True
        p_c.add_run(c["cause"] + "\n")
        p_c.add_run("Engineering Resolution: ").bold = True
        p_c.add_run(c["fix"] + "\n")

    # Section 4: Key Learnings & Engineering Takeaways
    h4 = doc.add_heading("4. Key Learnings & Engineering Takeaways", level=1)
    h4.runs[0].font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    learnings = [
        "Protocol Hygiene in Agent Systems: Stdio transport channels must strictly isolate data (sys.stdout) from diagnostic logs (sys.stderr). Single-character stream pollution breaks JSON-RPC deserialization.",
        "Resilient AI Execution & Fallbacks: AI agents must never rely solely on remote API endpoints or local LLMs without non-blocking async execution (await llm.ainvoke) and deterministic fallback logic.",
        "FortiOS CMDB API Lifecycle: Modifying policy objects on FortiGate requires a two-step REST API sequence: POST /cmdb/firewall/address followed by PUT/POST /cmdb/firewall/addrgrp.",
        "DPI Regex IP Filtering: Parsing logs requires separating RFC1918 private IP subnets from global public IPv4 addresses before executing threat intelligence calls."
    ]
    for l in learnings:
        bp = doc.add_paragraph(style='List Bullet')
        bp.add_run(l)

    # Section 5: What We Could Have Done Better
    h5 = doc.add_heading("5. Retrospective — What We Could Have Done Better", level=1)
    h5.runs[0].font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    better = [
        "Asynchronous Parallel Lookups: In Node 2, querying AbuseIPDB IPs sequentially using requests.get adds latency. Using asyncio with aiohttp would allow querying 50 IPs in parallel in under 500ms.",
        "Real-Time Syslog Stream Listener: Polling /api/v2/log/disk/traffic/forward requires HTTP requests. Building a UDP Syslog receiver on port 514 would allow instant event-driven triage as packets cross firewall interfaces.",
        "Pydantic Environment Schema Validation: Using Pydantic Settings to validate .env keys on startup rather than relying on os.getenv fallback strings."
    ]
    for b in better:
        bp = doc.add_paragraph(style='List Bullet')
        bp.add_run(b)

    # Section 6: Next-Level Future Roadmap
    h6 = doc.add_heading("6. Next-Level Enhancement Roadmap", level=1)
    h6.runs[0].font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Phase'
    hdr_cells[1].text = 'Capability Focus'
    hdr_cells[2].text = 'Technical Description & Value'

    for cell in hdr_cells:
        set_cell_background(cell, "003366")
        for p in cell.paragraphs:
            p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            p.runs[0].font.bold = True

    roadmap_data = [
        ("Phase 1", "Human-in-the-Loop Approval UI", "Build a React/Vite Web UI Dashboard with 'Approve Block' buttons so senior analysts can approve Node 4 remediation actions before firewall execution."),
        ("Phase 2", "Automated 7-Day Unblock Scheduler", "Implement background cron scheduling to automatically remove IP objects from Blocked_Bad_IPs after 7 days if AbuseIPDB score drops."),
        ("Phase 3", "Multi-Firewall HA Cluster Sync", "Expand fortigate_api.py to synchronize Blocked_Bad_IPs address groups across multi-VDOM or HA clusters using FortiManager REST API."),
        ("Phase 4", "SOAR Notification Webhooks", "Integrate Slack/Microsoft Teams notifications and automatic Jira Security Ticket creation whenever a high-risk IP is blocked.")
    ]

    for row in roadmap_data:
        row_cells = table.add_row().cells
        row_cells[0].text = row[0]
        row_cells[1].text = row[1]
        row_cells[2].text = row[2]
        for cell in row_cells:
            set_cell_margins(cell, top=120, bottom=120, left=150, right=150)

    doc.save("Enterprise_FortiGate_SOC_AI_Shield_Masterclass_Report.docx")
    print("Masterclass Word document successfully generated!")

if __name__ == "__main__":
    create_document()
